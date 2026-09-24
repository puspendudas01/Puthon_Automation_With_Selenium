"""Helpers to read test data from JSON and Excel files."""
import json
from pathlib import Path
from openpyxl import load_workbook


def read_json(path) -> dict:
    with open(Path(path), encoding="utf-8") as fh:
        return json.load(fh)


def read_excel(path, sheet: str) -> list[dict]:
    """Return each data row of `sheet` as a dict keyed by the header row."""
    wb = load_workbook(Path(path), data_only=True)
    ws = wb[sheet]
    rows = list(ws.iter_rows(values_only=True))
    headers = [str(h).strip() for h in rows[0]]
    return [dict(zip(headers, r)) for r in rows[1:] if any(c is not None for c in r)]
